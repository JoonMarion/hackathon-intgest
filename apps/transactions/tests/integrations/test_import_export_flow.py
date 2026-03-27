import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.categories.models import Category
from apps.core.tests.base import BaseIntegrationTestCase
from apps.core.tests.factories import CategoryFactory, FinancialAccountFactory, TransactionFactory, UserFactory
from apps.transactions.models import ExportJob, ImportJob, Transaction


class TransactionImportExportFlowIntegrationTests(BaseIntegrationTestCase):
    def setUp(self):
        self.user = UserFactory(username="transactions-import-export-user", email="transactions-import-export@example.com")
        self.client.force_login(self.user)
        self.expense_category = CategoryFactory(user=self.user, name="Alimentação", kind=Category.Kind.EXPENSE)
        self.income_category = CategoryFactory(user=self.user, name="Salário", kind=Category.Kind.INCOME)
        self.primary_account = FinancialAccountFactory(user=self.user, name="Conta principal")

    def _post_default_mapping(self, import_job):
        return self.client.post(
            reverse("transactions_http:import_mapping", args=[import_job.pk]),
            data={
                "date": "Date",
                "amount": "Amount",
                "payee": "Payee",
                "account": "Account",
                "import_id": "ImportId",
                "description": "Description",
                "notes": "Notes",
                "category": "Category",
            },
        )

    def test_import_csv_flow_processes_preview_and_confirmation(self):
        csv_payload = "".join(
            [
                "Date,Amount,Payee,Account,Description,Notes,Category,ImportId\n",
                "2026-03-20,-50.25,Mercado XYZ,Carteira,Supermercado,Feira mensal,Alimentação,row-1\n",
                "2026-03-25,3000.00,Empresa ABC,Conta Salário,Pagamento mensal,,Salário,row-2\n",
            ]
        )
        upload = SimpleUploadedFile("lote.csv", csv_payload.encode("utf-8"), content_type="text/csv")

        upload_response = self.client.post(reverse("transactions_http:import_upload"), data={"source_file": upload})
        import_job = ImportJob.objects.get(user=self.user)
        self.assertRedirects(upload_response, reverse("transactions_http:import_mapping", args=[import_job.pk]))

        mapping_response = self._post_default_mapping(import_job)
        self.assertRedirects(mapping_response, reverse("transactions_http:import_preview", args=[import_job.pk]))

        with self.captureOnCommitCallbacks(execute=True):
            confirm_response = self.client.post(reverse("transactions_http:import_confirm", args=[import_job.pk]))
        self.assertRedirects(confirm_response, reverse("transactions_http:import_status", args=[import_job.pk]))

        import_job.refresh_from_db()
        self.assertEqual(import_job.status, ImportJob.Status.COMPLETED)
        self.assertEqual(import_job.imported_rows, 2)
        self.assertEqual(import_job.skipped_rows, 0)
        self.assertEqual(import_job.failed_rows, 0)

        imported_expense = Transaction.objects.get(user=self.user, import_id="row-1")
        imported_income = Transaction.objects.get(user=self.user, import_id="row-2")
        self.assertEqual(imported_expense.kind, Transaction.Kind.EXPENSE)
        self.assertEqual(str(imported_expense.amount), "50.25")
        self.assertEqual(imported_expense.account.name, "Carteira")
        self.assertEqual(imported_income.kind, Transaction.Kind.INCOME)
        self.assertEqual(imported_income.account.name, "Conta Salário")

    def test_import_xlsx_flow_skips_existing_import_id(self):
        TransactionFactory(
            user=self.user,
            account=self.primary_account,
            category=self.expense_category,
            kind=Transaction.Kind.EXPENSE,
            amount="99.00",
            import_id="dup-1",
        )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Date", "Amount", "Payee", "Account", "Description", "Notes", "Category", "ImportId"])
        worksheet.append(["2026-03-26", -99, "Mercado XYZ", "Carteira", "Compra", "", "Alimentação", "dup-1"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        upload = SimpleUploadedFile(
            "lote.xlsx",
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        upload_response = self.client.post(reverse("transactions_http:import_upload"), data={"source_file": upload})
        import_job = ImportJob.objects.get(user=self.user)
        self.assertRedirects(upload_response, reverse("transactions_http:import_mapping", args=[import_job.pk]))

        mapping_response = self._post_default_mapping(import_job)
        self.assertRedirects(mapping_response, reverse("transactions_http:import_preview", args=[import_job.pk]))

        with self.captureOnCommitCallbacks(execute=True):
            confirm_response = self.client.post(reverse("transactions_http:import_confirm", args=[import_job.pk]))
        self.assertRedirects(confirm_response, reverse("transactions_http:import_status", args=[import_job.pk]))

        import_job.refresh_from_db()
        self.assertEqual(import_job.status, ImportJob.Status.COMPLETED)
        self.assertEqual(import_job.imported_rows, 0)
        self.assertEqual(import_job.skipped_rows, 1)
        self.assertEqual(import_job.failed_rows, 0)

    def test_export_csv_download_returns_generated_file(self):
        TransactionFactory(
            user=self.user,
            account=self.primary_account,
            category=self.expense_category,
            kind=Transaction.Kind.EXPENSE,
            amount="65.40",
            payee="Mercado XYZ",
            description="Compras",
        )
        TransactionFactory(
            user=self.user,
            account=self.primary_account,
            category=self.income_category,
            kind=Transaction.Kind.INCOME,
            amount="4200.00",
            payee="Empresa ABC",
            description="Salário",
        )

        with self.captureOnCommitCallbacks(execute=True):
            create_response = self.client.post(
                reverse("transactions_http:export_create"),
                data={
                    "format": "csv",
                    "q": "",
                    "category": "",
                    "kind": "",
                    "date_from": "",
                    "date_to": "",
                    "ordering": "-transaction_date,-created_at",
                },
            )

        export_job = ExportJob.objects.get(user=self.user, export_format=ExportJob.Format.CSV)
        self.assertRedirects(create_response, reverse("transactions_http:export_status", args=[export_job.pk]))

        export_job.refresh_from_db()
        self.assertEqual(export_job.status, ExportJob.Status.COMPLETED)
        self.assertEqual(export_job.row_count, 2)
        self.assertTrue(bool(export_job.output_file.file.name))

        download_response = self.client.get(reverse("transactions_http:export_download", args=[export_job.pk]))
        self.assertEqual(download_response.status_code, 200)
        payload = b"".join(download_response.streaming_content).decode("utf-8")
        self.assertIn("Date,Amount,Kind,Payee,Description,Account,Category,Notes,ImportId", payload)
        self.assertIn("Mercado XYZ", payload)
        self.assertIn("Empresa ABC", payload)

    def test_export_xlsx_download_returns_generated_file(self):
        TransactionFactory(
            user=self.user,
            account=self.primary_account,
            category=self.expense_category,
            kind=Transaction.Kind.EXPENSE,
            amount="10.50",
            payee="Padaria",
            description="Café",
        )

        with self.captureOnCommitCallbacks(execute=True):
            create_response = self.client.post(
                reverse("transactions_http:export_create"),
                data={
                    "format": "xlsx",
                    "q": "",
                    "category": "",
                    "kind": "",
                    "date_from": "",
                    "date_to": "",
                    "ordering": "-transaction_date,-created_at",
                },
            )

        export_job = ExportJob.objects.get(user=self.user, export_format=ExportJob.Format.XLSX)
        self.assertRedirects(create_response, reverse("transactions_http:export_status", args=[export_job.pk]))

        export_job.refresh_from_db()
        self.assertEqual(export_job.status, ExportJob.Status.COMPLETED)
        self.assertEqual(export_job.row_count, 1)

        download_response = self.client.get(reverse("transactions_http:export_download", args=[export_job.pk]))
        self.assertEqual(download_response.status_code, 200)

        payload = b"".join(download_response.streaming_content)
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))

        self.assertEqual(rows[0], ("Date", "Amount", "Kind", "Payee", "Description", "Account", "Category", "Notes", "ImportId"))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][3], "Padaria")
