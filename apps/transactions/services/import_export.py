import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.files.base import ContentFile
from django.db import IntegrityError
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.categories.models import Category
from apps.transactions.http.filters import TransactionFilterSet
from apps.transactions.models import DocumentFile, ExportJob, FinancialAccount, ImportJob, ImportRowResult, Transaction
from apps.transactions.services.accounts import ensure_default_account

REQUIRED_MAPPING_KEYS = ("date", "amount", "payee", "account")
OPTIONAL_MAPPING_KEYS = ("description", "notes", "category", "import_id")
ALLOWED_MAPPING_KEYS = REQUIRED_MAPPING_KEYS + OPTIONAL_MAPPING_KEYS

DEFAULT_ORDERING_VALUE = "-transaction_date,-created_at"
ORDERING_OPTIONS = {
    "-transaction_date,-created_at": ("-transaction_date", "-created_at"),
    "transaction_date,created_at": ("transaction_date", "created_at"),
    "-amount,-created_at": ("-amount", "-created_at"),
    "amount,created_at": ("amount", "created_at"),
}

EXPORT_HEADERS = [
    "Date",
    "Amount",
    "Kind",
    "Payee",
    "Description",
    "Account",
    "Category",
    "Notes",
    "ImportId",
]


def extract_headers(document_file: DocumentFile) -> list[str]:
    headers, _ = _load_tabular_rows(document_file)
    return headers


def build_import_preview(*, import_job: ImportJob, mapping: dict[str, str]) -> tuple[dict[str, int], list[dict[str, str]]]:
    clean_mapping = _normalize_mapping(mapping)
    _validate_mapping(clean_mapping)

    _, raw_rows = _load_tabular_rows(import_job.source_file)
    preview_rows: list[dict[str, str]] = []
    summary = {"total": 0, "importable": 0, "skipped": 0, "failed": 0}

    for row_number, raw_row in enumerate(raw_rows, start=2):
        summary["total"] += 1
        parsed = _parse_row(
            user=import_job.user,
            raw_row=raw_row,
            row_number=row_number,
            mapping=clean_mapping,
            dry_run=True,
        )
        preview_rows.append(
            {
                "row_number": str(row_number),
                "status": parsed["status"],
                "message": parsed["message"],
                "raw": _compact_raw_row(raw_row),
            }
        )
        if parsed["status"] == "ready":
            summary["importable"] += 1
        elif parsed["status"] == "skipped":
            summary["skipped"] += 1
        else:
            summary["failed"] += 1

    return summary, preview_rows


def run_import_job(*, import_job_id: int) -> dict[str, int | str]:
    import_job = ImportJob.objects.select_related("source_file").get(pk=import_job_id)
    import_job.status = ImportJob.Status.PROCESSING
    import_job.error_message = ""
    import_job.save(update_fields=["status", "error_message", "updated_at"])

    try:
        clean_mapping = _normalize_mapping(import_job.column_mapping)
        _validate_mapping(clean_mapping)

        import_job.rows.all().delete()

        _, raw_rows = _load_tabular_rows(import_job.source_file)
    except (ValueError, OSError) as exc:  # pragma: no cover - fallback defensivo
        import_job.status = ImportJob.Status.FAILED
        import_job.error_message = str(exc)
        import_job.save(update_fields=["status", "error_message", "updated_at"])
        return {"status": import_job.status, "imported": 0, "skipped": 0, "failed": 0}

    imported = 0
    skipped = 0
    failed = 0

    for row_number, raw_row in enumerate(raw_rows, start=2):
        parsed = _parse_row(
            user=import_job.user,
            raw_row=raw_row,
            row_number=row_number,
            mapping=clean_mapping,
            dry_run=False,
        )

        if parsed["status"] == "failed":
            failed += 1
            ImportRowResult.objects.create(
                import_job=import_job,
                row_number=row_number,
                status=ImportRowResult.Status.FAILED,
                message=parsed["message"],
                raw_data=_compact_raw_row(raw_row),
            )
            continue

        if parsed["status"] == "skipped":
            skipped += 1
            ImportRowResult.objects.create(
                import_job=import_job,
                row_number=row_number,
                status=ImportRowResult.Status.SKIPPED,
                message=parsed["message"],
                raw_data=_compact_raw_row(raw_row),
            )
            continue

        payload = parsed["payload"]
        account = _resolve_account(user=import_job.user, name=payload["account_name"])
        category = _resolve_category(
            user=import_job.user,
            kind=payload["kind"],
            category_name=payload["category_name"],
        )

        try:
            transaction = Transaction.objects.create(
                user=import_job.user,
                account=account,
                category=category,
                kind=payload["kind"],
                amount=payload["amount"],
                transaction_date=payload["transaction_date"],
                payee=payload["payee"],
                description=payload["description"],
                notes=payload["notes"],
                import_id=payload["import_id"],
            )
        except IntegrityError:
            skipped += 1
            ImportRowResult.objects.create(
                import_job=import_job,
                row_number=row_number,
                status=ImportRowResult.Status.SKIPPED,
                message="Linha já importada por ID de importação.",
                raw_data=_compact_raw_row(raw_row),
            )
            continue

        imported += 1
        ImportRowResult.objects.create(
            import_job=import_job,
            row_number=row_number,
            status=ImportRowResult.Status.IMPORTED,
            message="Importado com sucesso.",
            raw_data=_compact_raw_row(raw_row),
            transaction=transaction,
        )

    import_job.status = ImportJob.Status.COMPLETED
    import_job.imported_rows = imported
    import_job.skipped_rows = skipped
    import_job.failed_rows = failed
    import_job.save(
        update_fields=[
            "status",
            "imported_rows",
            "skipped_rows",
            "failed_rows",
            "updated_at",
        ]
    )

    return {
        "status": import_job.status,
        "imported": imported,
        "skipped": skipped,
        "failed": failed,
    }


def run_export_job(*, export_job_id: int) -> dict[str, int | str]:
    export_job = ExportJob.objects.select_related("output_file", "user").get(pk=export_job_id)
    export_job.status = ExportJob.Status.PROCESSING
    export_job.error_message = ""
    export_job.save(update_fields=["status", "error_message", "updated_at"])

    try:
        queryset = build_export_queryset(user=export_job.user, filters=export_job.filters)
        rows = [
            _serialize_transaction_row(transaction=transaction)
            for transaction in queryset
        ]

        if export_job.export_format == ExportJob.Format.CSV:
            payload = _render_csv(rows)
            extension = "csv"
            content_type = "text/csv"
        else:
            payload = _render_xlsx(rows)
            extension = "xlsx"
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        timestamp = timezone.localtime().strftime("%Y%m%d-%H%M%S")
        filename = f"transactions-export-{export_job.id}-{timestamp}.{extension}"

        export_job.output_file.file.save(filename, ContentFile(payload), save=False)
        export_job.output_file.content_type = content_type
        export_job.output_file.size_bytes = len(payload)
        export_job.output_file.original_name = filename
        export_job.output_file.save(update_fields=["file", "content_type", "size_bytes", "original_name", "updated_at"])

        export_job.status = ExportJob.Status.COMPLETED
        export_job.row_count = len(rows)
        export_job.save(update_fields=["status", "row_count", "updated_at"])

        return {
            "status": export_job.status,
            "row_count": len(rows),
            "export_format": export_job.export_format,
        }
    except (ValueError, OSError) as exc:  # pragma: no cover - fallback defensivo
        export_job.status = ExportJob.Status.FAILED
        export_job.error_message = str(exc)
        export_job.save(update_fields=["status", "error_message", "updated_at"])
        return {"status": export_job.status, "row_count": 0, "export_format": export_job.export_format}


def build_export_queryset(*, user, filters: dict[str, str]):
    base_queryset = Transaction.objects.filter(user=user).select_related("category", "account")
    filterset = TransactionFilterSet(data=filters, queryset=base_queryset, user=user)

    ordering_value = filters.get("ordering", DEFAULT_ORDERING_VALUE)
    ordering_fields = ORDERING_OPTIONS.get(ordering_value, ORDERING_OPTIONS[DEFAULT_ORDERING_VALUE])
    return filterset.qs.order_by(*ordering_fields)


def _load_tabular_rows(document_file: DocumentFile) -> tuple[list[str], list[dict[str, str]]]:
    extension = Path(document_file.original_name).suffix.lower()
    if extension == ".csv":
        return _load_csv_rows(document_file)
    if extension == ".xlsx":
        return _load_xlsx_rows(document_file)
    raise ValueError("Formato de arquivo não suportado.")


def _load_csv_rows(document_file: DocumentFile) -> tuple[list[str], list[dict[str, str]]]:
    document_file.file.open("rb")
    try:
        payload = document_file.file.read()
    finally:
        document_file.file.close()

    try:
        content = payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        content = payload.decode("latin-1")

    reader = csv.DictReader(io.StringIO(content))
    headers = [str(header).strip() for header in (reader.fieldnames or []) if header]
    rows: list[dict[str, str]] = []

    for row in reader:
        normalized_row: dict[str, str] = {}
        for header in headers:
            normalized_row[header] = (row.get(header) or "").strip()
        rows.append(normalized_row)

    return headers, rows


def _load_xlsx_rows(document_file: DocumentFile) -> tuple[list[str], list[dict[str, str]]]:
    document_file.file.open("rb")
    try:
        workbook = load_workbook(document_file.file, read_only=True, data_only=True)
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)

        raw_headers = next(iterator, None)
        if raw_headers is None:
            return [], []

        headers = [str(value).strip() if value is not None else "" for value in raw_headers]

        rows: list[dict[str, str]] = []
        for values in iterator:
            normalized_row: dict[str, str] = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = values[index] if index < len(values) else None
                normalized_row[header] = _stringify_cell(value)
            rows.append(normalized_row)

        workbook.close()
        return [header for header in headers if header], rows
    finally:
        document_file.file.close()


def _stringify_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _normalize_mapping(mapping: dict[str, str]) -> dict[str, str]:
    clean_mapping: dict[str, str] = {}
    for key in ALLOWED_MAPPING_KEYS:
        value = mapping.get(key, "")
        clean_mapping[key] = str(value).strip()
    return clean_mapping


def _validate_mapping(mapping: dict[str, str]) -> None:
    missing = [key for key in REQUIRED_MAPPING_KEYS if not mapping.get(key)]
    if missing:
        raise ValueError("Mapeamento incompleto para colunas obrigatórias.")


def _parse_row(*, user, raw_row: dict[str, str], row_number: int, mapping: dict[str, str], dry_run: bool) -> dict[str, object]:
    del row_number
    row_payload = {
        "transaction_date": _pick_cell(raw_row, mapping["date"]),
        "amount": _pick_cell(raw_row, mapping["amount"]),
        "payee": _pick_cell(raw_row, mapping["payee"]),
        "account_name": _pick_cell(raw_row, mapping["account"]),
        "description": _pick_cell(raw_row, mapping.get("description", "")),
        "notes": _pick_cell(raw_row, mapping.get("notes", "")),
        "category_name": _pick_cell(raw_row, mapping.get("category", "")),
        "import_id": _pick_cell(raw_row, mapping.get("import_id", "")),
    }

    if not row_payload["transaction_date"]:
        return {"status": "failed", "message": "Data obrigatória ausente."}
    if not row_payload["amount"]:
        return {"status": "failed", "message": "Valor obrigatório ausente."}
    if not row_payload["payee"]:
        return {"status": "failed", "message": "Favorecido obrigatório ausente."}
    if not row_payload["account_name"]:
        return {"status": "failed", "message": "Conta obrigatória ausente."}

    parsed_date = _parse_date(row_payload["transaction_date"])
    if parsed_date is None:
        return {"status": "failed", "message": "Data inválida."}

    parsed_amount = _parse_decimal(row_payload["amount"])
    if parsed_amount is None or parsed_amount == Decimal("0"):
        return {"status": "failed", "message": "Valor inválido."}

    kind = Transaction.Kind.INCOME if parsed_amount > 0 else Transaction.Kind.EXPENSE
    normalized_amount = abs(parsed_amount).quantize(Decimal("0.01"))

    import_id = row_payload["import_id"].strip()
    if import_id:
        if Transaction.objects.filter(user=user, import_id=import_id).exists():
            return {"status": "skipped", "message": "Duplicado por ImportId."}
    else:
        payee = row_payload["payee"].strip()
        if Transaction.objects.filter(
            user=user,
            transaction_date=parsed_date,
            amount=normalized_amount,
            payee__iexact=payee,
        ).exists():
            return {"status": "skipped", "message": "Duplicado por data+valor+favorecido."}

    payee = row_payload["payee"].strip()

    payload = {
        "transaction_date": parsed_date,
        "amount": normalized_amount,
        "kind": kind,
        "payee": payee,
        "account_name": row_payload["account_name"].strip(),
        "description": row_payload["description"].strip(),
        "notes": row_payload["notes"].strip(),
        "category_name": row_payload["category_name"].strip(),
        "import_id": import_id,
    }

    if dry_run:
        return {"status": "ready", "message": "Pronto para importação.", "payload": payload}
    return {"status": "ready", "message": "Pronto para importar.", "payload": payload}


def _pick_cell(row: dict[str, str], column_name: str) -> str:
    if not column_name:
        return ""
    return str(row.get(column_name, "")).strip()


def _parse_date(raw_value: str) -> date | None:
    if not raw_value:
        return None

    candidates = ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d"]
    for fmt in candidates:
        try:
            return datetime.strptime(raw_value, fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(raw_value)
    except ValueError:
        return None


def _parse_decimal(raw_value: str) -> Decimal | None:
    value = raw_value.strip().replace("R$", "").replace(" ", "")
    if not value:
        return None

    if "," in value and "." in value:
        if value.rfind(",") > value.rfind("."):
            value = value.replace(".", "").replace(",", ".")
        else:
            value = value.replace(",", "")
    elif "," in value:
        value = value.replace(".", "").replace(",", ".")

    try:
        return Decimal(value)
    except (InvalidOperation, ValueError):
        return None


def _resolve_account(*, user, name: str) -> FinancialAccount:
    account_name = name.strip() or "Conta principal"
    account, _ = FinancialAccount.objects.get_or_create(
        user=user,
        name=account_name,
        defaults={"is_active": True},
    )
    if not account.is_active:
        account.is_active = True
        account.save(update_fields=["is_active", "updated_at"])
    return account


def _resolve_category(*, user, kind: str, category_name: str) -> Category:
    fallback_name = "Importado - Receita" if kind == Transaction.Kind.INCOME else "Importado - Despesa"
    name = category_name.strip() or fallback_name

    category = Category.objects.filter(user=user, name__iexact=name, kind=kind).first()
    if category:
        return category

    conflicting = Category.objects.filter(user=user, name__iexact=name).first()
    if conflicting is not None and conflicting.kind != kind:
        name = fallback_name

    category, _ = Category.objects.get_or_create(
        user=user,
        name=name,
        defaults={"kind": kind},
    )
    if category.kind != kind:
        category.name = fallback_name
        category.kind = kind
        category.save(update_fields=["name", "kind", "updated_at"])
    return category


def _compact_raw_row(raw_row: dict[str, str]) -> dict[str, str]:
    return {key: str(value) for key, value in raw_row.items()}


def _serialize_transaction_row(*, transaction: Transaction) -> list[str]:
    signed_amount = transaction.amount if transaction.kind == Transaction.Kind.INCOME else -transaction.amount
    return [
        transaction.transaction_date.isoformat(),
        str(signed_amount),
        transaction.kind,
        transaction.payee,
        transaction.description,
        transaction.account.name,
        transaction.category.name,
        transaction.notes,
        transaction.import_id,
    ]


def _render_csv(rows: list[list[str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADERS)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def _render_xlsx(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Transactions"
    worksheet.append(EXPORT_HEADERS)

    for row in rows:
        worksheet.append(row)

    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def ensure_user_minimum_account(*, user) -> FinancialAccount:
    account = FinancialAccount.objects.filter(user=user).order_by("name").first()
    if account:
        return account
    return ensure_default_account(user=user)
